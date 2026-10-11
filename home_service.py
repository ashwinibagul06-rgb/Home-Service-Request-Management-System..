import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "service_requests.xlsx"


# ---------------- EXCEL FILE ----------------

def create_excel():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Service Requests"

        ws.append([
            "ID",
            "Customer Name",
            "Mobile",
            "Address",
            "Service Type",
            "Problem",
            "Date",
            "Status"
        ])

        wb.save(FILE_NAME)


# ---------------- LOGIN ----------------

def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "1234":
        login_window.destroy()
        dashboard()
    else:
        messagebox.showerror(
            "Login Failed",
            "Invalid username or password"
        )


# ---------------- DASHBOARD ----------------

def dashboard():

    global root

    root = tk.Tk()
    root.title("Home Service Request Management System")
    root.geometry("900x600")

    title = tk.Label(
        root,
        text="HOME SERVICE REQUEST MANAGEMENT SYSTEM",
        font=("Arial", 18, "bold")
    )
    title.pack(pady=20)

    button_frame = tk.Frame(root)
    button_frame.pack(pady=20)

    tk.Button(
        button_frame,
        text="Add Request",
        width=18,
        command=add_request
    ).grid(row=0, column=0, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="View Requests",
        width=18,
        command=view_requests
    ).grid(row=0, column=1, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="Update Request",
        width=18,
        command=update_request
    ).grid(row=1, column=0, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="Delete Request",
        width=18,
        command=delete_request
    ).grid(row=1, column=1, padx=10, pady=10)

    tk.Button(
        root,
        text="Logout",
        width=20,
        command=logout
    ).pack(pady=30)

    root.mainloop()


# ---------------- ADD REQUEST ----------------

def add_request():

    window = tk.Toplevel(root)
    window.title("Add Service Request")
    window.geometry("500x600")

    tk.Label(
        window,
        text="Add Service Request",
        font=("Arial", 16, "bold")
    ).pack(pady=15)

    form = tk.Frame(window)
    form.pack()

    tk.Label(form, text="Customer Name").grid(row=0, column=0, pady=8)
    name_entry = tk.Entry(form, width=30)
    name_entry.grid(row=0, column=1)

    tk.Label(form, text="Mobile").grid(row=1, column=0, pady=8)
    mobile_entry = tk.Entry(form, width=30)
    mobile_entry.grid(row=1, column=1)

    tk.Label(form, text="Address").grid(row=2, column=0, pady=8)
    address_entry = tk.Entry(form, width=30)
    address_entry.grid(row=2, column=1)

    tk.Label(form, text="Service Type").grid(row=3, column=0, pady=8)

    service_box = ttk.Combobox(
        form,
        values=[
            "Plumbing",
            "Electrical",
            "Cleaning",
            "AC Repair",
            "Carpenter",
            "Other"
        ],
        width=27
    )
    service_box.grid(row=3, column=1)
    service_box.current(0)

    tk.Label(form, text="Problem").grid(row=4, column=0, pady=8)
    problem_entry = tk.Entry(form, width=30)
    problem_entry.grid(row=4, column=1)

    tk.Label(form, text="Date").grid(row=5, column=0, pady=8)
    date_entry = tk.Entry(form, width=30)
    date_entry.grid(row=5, column=1)

    tk.Label(form, text="Status").grid(row=6, column=0, pady=8)

    status_box = ttk.Combobox(
        form,
        values=[
            "Pending",
            "In Progress",
            "Completed"
        ],
        width=27
    )
    status_box.grid(row=6, column=1)
    status_box.current(0)

    def save_request():

        name = name_entry.get()
        mobile = mobile_entry.get()
        address = address_entry.get()
        service = service_box.get()
        problem = problem_entry.get()
        date = date_entry.get()
        status = status_box.get()

        if name == "" or mobile == "" or address == "" or problem == "":
            messagebox.showwarning(
                "Warning",
                "Please fill all required fields"
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        new_id = ws.max_row

        ws.append([
            new_id,
            name,
            mobile,
            address,
            service,
            problem,
            date,
            status
        ])

        wb.save(FILE_NAME)

        messagebox.showinfo(
            "Success",
            "Service request saved successfully!"
        )

        window.destroy()

    tk.Button(
        window,
        text="Save Request",
        width=20,
        command=save_request
    ).pack(pady=20)

    tk.Button(
        window,
        text="Clear",
        width=20,
        command=lambda: clear_fields(
            name_entry,
            mobile_entry,
            address_entry,
            problem_entry,
            date_entry
        )
    ).pack()


# ---------------- CLEAR ----------------

def clear_fields(*entries):

    for entry in entries:
        entry.delete(0, tk.END)


# ---------------- VIEW REQUESTS ----------------

def view_requests():

    window = tk.Toplevel(root)
    window.title("View Service Requests")
    window.geometry("1000x500")

    tk.Label(
        window,
        text="Service Requests",
        font=("Arial", 16, "bold")
    ).pack(pady=10)

    columns = (
        "ID",
        "Customer",
        "Mobile",
        "Address",
        "Service",
        "Problem",
        "Date",
        "Status"
    )

    tree = ttk.Treeview(
        window,
        columns=columns,
        show="headings"
    )

    for column in columns:
        tree.heading(column, text=column)
        tree.column(column, width=110)

    tree.pack(fill="both", expand=True)

    wb = load_workbook(FILE_NAME)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)


# ---------------- UPDATE REQUEST ----------------

def update_request():

    window = tk.Toplevel(root)
    window.title("Update Service Request")
    window.geometry("500x600")

    tk.Label(
        window,
        text="Update Service Request",
        font=("Arial", 16, "bold")
    ).pack(pady=15)

    form = tk.Frame(window)
    form.pack()

    tk.Label(form, text="Request ID").grid(row=0, column=0, pady=8)

    id_entry = tk.Entry(form, width=30)
    id_entry.grid(row=0, column=1)

    tk.Label(form, text="Customer Name").grid(row=1, column=0, pady=8)

    name_entry = tk.Entry(form, width=30)
    name_entry.grid(row=1, column=1)

    tk.Label(form, text="Mobile").grid(row=2, column=0, pady=8)

    mobile_entry = tk.Entry(form, width=30)
    mobile_entry.grid(row=2, column=1)

    tk.Label(form, text="Address").grid(row=3, column=0, pady=8)

    address_entry = tk.Entry(form, width=30)
    address_entry.grid(row=3, column=1)

    tk.Label(form, text="Service Type").grid(row=4, column=0, pady=8)

    service_box = ttk.Combobox(
        form,
        values=[
            "Plumbing",
            "Electrical",
            "Cleaning",
            "AC Repair",
            "Carpenter",
            "Other"
        ],
        width=27
    )
    service_box.grid(row=4, column=1)

    tk.Label(form, text="Problem").grid(row=5, column=0, pady=8)

    problem_entry = tk.Entry(form, width=30)
    problem_entry.grid(row=5, column=1)

    tk.Label(form, text="Date").grid(row=6, column=0, pady=8)

    date_entry = tk.Entry(form, width=30)
    date_entry.grid(row=6, column=1)

    tk.Label(form, text="Status").grid(row=7, column=0, pady=8)

    status_box = ttk.Combobox(
        form,
        values=[
            "Pending",
            "In Progress",
            "Completed"
        ],
        width=27
    )
    status_box.grid(row=7, column=1)

    def load_data():

        request_id = id_entry.get()

        if request_id == "":
            messagebox.showwarning(
                "Warning",
                "Enter Request ID"
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        found = False

        for row in ws.iter_rows(min_row=2):

            if str(row[0].value) == request_id:

                name_entry.delete(0, tk.END)
                name_entry.insert(0, row[1].value)

                mobile_entry.delete(0, tk.END)
                mobile_entry.insert(0, row[2].value)

                address_entry.delete(0, tk.END)
                address_entry.insert(0, row[3].value)

                service_box.set(row[4].value)

                problem_entry.delete(0, tk.END)
                problem_entry.insert(0, row[5].value)

                date_entry.delete(0, tk.END)
                date_entry.insert(0, row[6].value)

                status_box.set(row[7].value)

                found = True
                break

        if not found:
            messagebox.showerror(
                "Error",
                "Request ID not found"
            )

    def save_update():

        request_id = id_entry.get()

        if request_id == "":
            messagebox.showwarning(
                "Warning",
                "Enter Request ID"
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        found = False

        for row in ws.iter_rows(min_row=2):

            if str(row[0].value) == request_id:

                row[1].value = name_entry.get()
                row[2].value = mobile_entry.get()
                row[3].value = address_entry.get()
                row[4].value = service_box.get()
                row[5].value = problem_entry.get()
                row[6].value = date_entry.get()
                row[7].value = status_box.get()

                found = True
                break

        if found:

            wb.save(FILE_NAME)

            messagebox.showinfo(
                "Success",
                "Request updated successfully!"
            )

            window.destroy()

        else:

            messagebox.showerror(
                "Error",
                "Request ID not found"
            )

    tk.Button(
        window,
        text="Load Request",
        width=20,
        command=load_data
    ).pack(pady=10)

    tk.Button(
        window,
        text="Update Request",
        width=20,
        command=save_update
    ).pack(pady=10)


# ---------------- DELETE REQUEST ----------------

def delete_request():

    window = tk.Toplevel(root)
    window.title("Delete Service Request")
    window.geometry("400x250")

    tk.Label(
        window,
        text="Delete Service Request",
        font=("Arial", 16, "bold")
    ).pack(pady=20)

    tk.Label(
        window,
        text="Enter Request ID:"
    ).pack()

    id_entry = tk.Entry(window, width=30)
    id_entry.pack(pady=10)

    def delete():

        request_id = id_entry.get()

        if request_id == "":
            messagebox.showwarning(
                "Warning",
                "Enter Request ID"
            )
            return

        answer = messagebox.askyesno(
            "Confirm Delete",
            "Are you sure you want to delete this request?"
        )

        if not answer:
            return

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        found = False

        for row in range(2, ws.max_row + 1):

            if str(ws.cell(row, 1).value) == request_id:

                ws.delete_rows(row, 1)
                found = True
                break

        if found:

            wb.save(FILE_NAME)

            messagebox.showinfo(
                "Success",
                "Request deleted successfully!"
            )

            window.destroy()

        else:

            messagebox.showerror(
                "Error",
                "Request ID not found"
            )

    tk.Button(
        window,
        text="Delete",
        width=20,
        command=delete
    ).pack(pady=10)


# ---------------- LOGOUT ----------------

def logout():

    answer = messagebox.askyesno(
        "Logout",
        "Do you want to logout?"
    )

    if answer:

        root.destroy()
        login_page()


# ---------------- LOGIN PAGE ----------------

def login_page():

    global login_window
    global username_entry
    global password_entry

    login_window = tk.Tk()
    login_window.title("Login")
    login_window.geometry("400x350")

    tk.Label(
        login_window,
        text="HOME SERVICE",
        font=("Arial", 20, "bold")
    ).pack(pady=20)

    tk.Label(
        login_window,
        text="Request Management System",
        font=("Arial", 12)
    ).pack(pady=5)

    tk.Label(
        login_window,
        text="Username"
    ).pack(pady=10)

    username_entry = tk.Entry(
        login_window,
        width=30
    )
    username_entry.pack()

    tk.Label(
        login_window,
        text="Password"
    ).pack(pady=10)

    password_entry = tk.Entry(
        login_window,
        width=30,
        show="*"
    )
    password_entry.pack()

    tk.Button(
        login_window,
        text="Login",
        width=20,
        command=login
    ).pack(pady=25)

    tk.Label(
        login_window,
        text="Username: admin\nPassword: 1234"
    ).pack()

    login_window.mainloop()


# ---------------- START PROGRAM ----------------

create_excel()
login_page()